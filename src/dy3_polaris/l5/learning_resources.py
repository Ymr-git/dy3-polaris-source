"""Truth-preserving personalized learning resource projection.

Resources are assembled from an existing Diagnosis-owned teaching decision,
R06 Concept context, a reviewed collaboration result, and the existing
PracticeBank entry point.  The module does not generate scientific facts,
experimental optima, questions, mastery, or a second learner state.
"""
from __future__ import annotations

from dataclasses import dataclass
from enum import Enum
from hashlib import sha256
import time
from typing import Any, Mapping

from dy3_polaris.l5.agent_contracts import (
    FinalCollaborationResult,
    QualityReleaseDecision,
    QualityReleaseStatus,
)
from dy3_polaris.l5.knowledge_learning_fusion import KnowledgeLearningContext
from dy3_polaris.l5.learner_foundation import AdaptiveTeachingDecision


class ResourceFamily(str, Enum):
    KNOWLEDGE = "knowledge_understanding"
    PRACTICAL = "research_practice"
    ASSESSMENT = "assessment_practice"


class ResourceSourceType(str, Enum):
    RETRIEVED = "retrieved"
    GENERATED = "generated"
    DERIVED = "derived"
    TEMPLATE = "template"


class ResourceInteractionAction(str, Enum):
    OPEN = "open"
    UNDERSTOOD = "understood"
    STILL_CONFUSED = "still_confused"
    CHANGE_EXPLANATION = "change_explanation"
    REQUEST_EXAMPLE = "request_example"
    DEEPEN = "deepen"
    ASK_FOLLOW_UP = "ask_follow_up"
    START_PRACTICE = "start_practice"
    NEXT_CONCEPT = "next_concept"


@dataclass(frozen=True, slots=True)
class LearningResource:
    resource_id: str
    resource_family: ResourceFamily
    resource_form: str
    title: str
    target_concepts: tuple[str, ...]
    prerequisite_concepts: tuple[str, ...]
    learning_goal: str
    learner_fit_reason: str
    difficulty: str
    estimated_time_minutes: int
    payload: Mapping[str, Any]
    source_type: ResourceSourceType
    evidence_refs: tuple[str, ...]
    review_status: str
    provenance: tuple[str, ...]
    interaction_actions: tuple[ResourceInteractionAction, ...]
    completion_signal: str
    next_action: str


@dataclass(frozen=True, slots=True)
class LearningResourcePlan:
    task_id: str
    learner_id: str
    teaching_depth: str
    strategy: str
    resources: tuple[LearningResource, ...]


@dataclass(frozen=True, slots=True)
class ResourceInteractionEvent:
    event_id: str
    learner_id: str
    task_id: str
    resource_id: str
    resource_family: str
    resource_form: str
    action: ResourceInteractionAction
    concept_ids: tuple[str, ...]
    source_type: str
    timestamp: float


def _resource_id(task_id: str, family: str, form: str) -> str:
    digest = sha256(f"{task_id}|{family}|{form}".encode("utf-8")).hexdigest()[:16]
    return f"resource-{digest}"


def _concept_names(
    context: KnowledgeLearningContext | None,
    concept_ids: tuple[str, ...],
) -> tuple[str, ...]:
    if context is None:
        return concept_ids
    return tuple(context.concept_names.get(item, item) for item in concept_ids)


def _reviewed_claims(
    final_result: FinalCollaborationResult | None,
) -> tuple[dict[str, Any], ...]:
    """Project only claims already accepted by the real collaboration result."""

    if not isinstance(final_result, FinalCollaborationResult):
        return ()
    return tuple(
        {
            "claim_id": str(claim.claim_id),
            "statement": str(claim.statement),
            "claim_type": str(getattr(claim.claim_type, "value", claim.claim_type)),
            "support_status": str(
                getattr(claim.support_status, "value", claim.support_status)
            ),
            "evidence_refs": tuple(str(item) for item in claim.evidence_refs if str(item)),
            "conditions": tuple(
                {"name": str(name), "value": str(value)}
                for name, value in claim.conditions
            ),
        }
        for claim in final_result.accepted_claims
        if str(claim.statement).strip()
    )


def _guided_questions(
    concept_names: tuple[str, ...],
    prerequisite_names: tuple[str, ...],
) -> tuple[dict[str, str], ...]:
    """Create prompts, never scientific answers, from real concept identities."""

    focus = concept_names[0] if concept_names else "当前核心概念"
    questions = [
        {
            "question_id": "self-explain",
            "prompt": f"请用自己的话解释“{focus}”的核心机制，并区分事实与推断。",
            "purpose": "SELF_EXPLANATION",
            "source_class": "RELATION_BACKED_PROMPT",
        },
        {
            "question_id": "evidence-boundary",
            "prompt": "当前证据能支持到什么程度？还有哪些条件或不确定性不能忽略？",
            "purpose": "EVIDENCE_BOUNDARY",
            "source_class": "RELATION_BACKED_PROMPT",
        },
    ]
    if prerequisite_names:
        questions.insert(0, {
            "question_id": "prerequisite-check",
            "prompt": f"在继续前，请先说明“{prerequisite_names[0]}”与当前问题的关系。",
            "purpose": "PREREQUISITE_CHECK",
            "source_class": "RELATION_BACKED_PROMPT",
        })
    return tuple(questions)


def _guided_document(
    *,
    goal: str,
    reviewed_answer: str,
    concept_names: tuple[str, ...],
    prerequisite_names: tuple[str, ...],
    accepted_claims: tuple[dict[str, Any], ...],
    knowledge_gaps: tuple[str, ...],
    evidence_refs: tuple[str, ...],
) -> dict[str, Any]:
    """Assemble a teaching-first reading view from released runtime facts.

    Evidence remains available as an appendix, but it is not used as a
    substitute for the lesson body.  The lesson body is the already-reviewed
    public answer; concept identities and accepted claims only organise that
    answer and never create additional scientific facts.
    """

    sections: list[dict[str, Any]] = [
        {
            "section_id": "learning-goal",
            "title": "学习目标",
            "source_class": "DECISION",
            "content": goal,
        },
        {
            "section_id": "reviewed-explanation",
            "title": "核心讲解",
            "source_class": "REVIEWED_OUTPUT",
            "content": reviewed_answer,
        },
    ]
    if accepted_claims:
        sections.append({
            "section_id": "accepted-claims",
            "title": "关键科学判断",
            "source_class": "REVIEWED_CLAIMS",
            "items": accepted_claims,
        })
    if concept_names or prerequisite_names:
        sections.append({
            "section_id": "concept-path",
            "title": "概念与先修关系",
            "source_class": "R06_CONCEPT_RELATION",
            "target_concepts": concept_names,
            "prerequisites": prerequisite_names,
        })
    sections.append({
        "section_id": "evidence-and-limits",
        "title": "证据与边界（附录）",
        "source_class": "RELEASE_GATE",
        "evidence_refs": evidence_refs,
        "knowledge_gaps": knowledge_gaps,
    })
    return {
        "document_mode": "GUIDED_LONG_READ",
        "lesson_sequence": (
            {"step": "ORIENT", "label": "明确目标"},
            {"step": "UNDERSTAND", "label": "理解讲解"},
            {"step": "CONNECT", "label": "连接概念"},
            {"step": "CHECK", "label": "检查理解"},
        ),
        "sections": tuple(sections),
        "fact_inference_separation": True,
        "source_policy": "released_runtime_facts_only",
    }


def build_learning_resource_plan(
    *,
    task_id: str,
    learner_id: str,
    teaching_decision: AdaptiveTeachingDecision | None,
    knowledge_context: KnowledgeLearningContext | None,
    final_result: FinalCollaborationResult | None,
    quality_release: QualityReleaseDecision,
    reviewed_long_form: Mapping[str, Any] | None = None,
    reviewed_guided_questions: Mapping[str, Any] | None = None,
) -> LearningResourcePlan:
    """Build three usable resource families without inventing content."""

    depth = (
        teaching_decision.content_depth
        if isinstance(teaching_decision, AdaptiveTeachingDecision)
        else "foundation"
    )
    strategy = (
        teaching_decision.explanation_strategy
        if isinstance(teaching_decision, AdaptiveTeachingDecision)
        else "baseline_explanation"
    )
    target_concepts = (
        tuple(knowledge_context.target_concepts)
        if isinstance(knowledge_context, KnowledgeLearningContext)
        else ()
    )
    prerequisite_concepts = (
        tuple(knowledge_context.learning_path.prerequisite_gap)
        if isinstance(knowledge_context, KnowledgeLearningContext)
        else ()
    )
    next_concept = (
        knowledge_context.learning_path.next_concept
        if isinstance(knowledge_context, KnowledgeLearningContext)
        else ""
    )
    concept_names = _concept_names(knowledge_context, target_concepts)
    prerequisite_names = _concept_names(knowledge_context, prerequisite_concepts)
    goal = str(
        getattr(getattr(final_result, "decision", None), "next_action", "")
        or "继续当前科研学习任务"
    )
    release_eligible = bool(
        quality_release.eligible
        and quality_release.status
        in {QualityReleaseStatus.FULL_RELEASE, QualityReleaseStatus.LIMITED_RELEASE}
    )
    # Resource and practice projections are part of the released teaching
    # result.  A withheld/failed review must not issue actionable follow-up
    # material derived from an unpublished answer.
    if not release_eligible:
        return LearningResourcePlan(
            task_id=task_id,
            learner_id=learner_id,
            teaching_depth=depth,
            strategy=strategy,
            resources=(),
        )
    evidence_refs = (
        tuple(
            str(item)
            for item in getattr(final_result, "provenance_refs", ()) or ()
            if str(item)
        )
        if release_eligible
        else ()
    )
    reviewed_answer = (
        quality_release.public_answer
        if quality_release.status in {
            QualityReleaseStatus.FULL_RELEASE,
            QualityReleaseStatus.LIMITED_RELEASE,
        }
        else ""
    )
    resource_modes = set(
        teaching_decision.resource_modes
        if isinstance(teaching_decision, AdaptiveTeachingDecision)
        else ()
    )
    accepted_claims = _reviewed_claims(final_result)
    knowledge_gaps = (
        tuple(str(item) for item in final_result.knowledge_gaps if str(item))
        if isinstance(final_result, FinalCollaborationResult)
        else ()
    )
    question_review = dict(reviewed_guided_questions or {})
    model_questions = tuple(
        {
            "question_id": str(item.get("question_id") or f"model-guided-{index}"),
            "prompt": str(item.get("prompt") or "").strip(),
            "purpose": str(item.get("purpose") or "SOCRATIC_MECHANISM"),
            "source_class": "GENERATION_REVIEWED_GUIDANCE",
        }
        for index, item in enumerate(question_review.get("questions") or (), 1)
        if isinstance(item, Mapping) and str(item.get("prompt") or "").strip()
    )
    model_questions_approved = bool(
        model_questions
        and bool(question_review.get("model_used", False))
        and bool(question_review.get("reviewer_executed", False))
        and str(question_review.get("review_verdict") or "").lower() == "approved"
    )
    guided_questions = (
        model_questions
        if model_questions_approved
        else _guided_questions(concept_names, prerequisite_names)
    )
    guided_document = _guided_document(
        goal=goal,
        reviewed_answer=reviewed_answer,
        concept_names=concept_names,
        prerequisite_names=prerequisite_names,
        accepted_claims=accepted_claims,
        knowledge_gaps=knowledge_gaps,
        evidence_refs=evidence_refs,
    )
    long_form = dict(reviewed_long_form or {})
    long_form_content = str(long_form.get("content") or "").strip()
    # Only a real model-generated resource may be presented as a generated
    # long-form lesson.  The deterministic evidence reader is useful for
    # internal inspection, but publishing it as a long lesson makes an
    # evidence appendix look like personalised teaching content.
    long_form_approved = bool(
        long_form_content
        and bool(long_form.get("model_used", False))
        and str(long_form.get("review_verdict") or "").lower() == "approved"
        and bool(long_form.get("reviewer_executed", False))
    )
    if long_form_approved:
        long_section = {
            "section_id": "reviewed-topic-long-form",
            "title": "Reviewer 通过的专题长文",
            "source_class": "GENERATION_REVIEWED_RESOURCE",
            "content": long_form_content,
        }
        current_sections = list(guided_document.get("sections") or ())
        current_sections.insert(2, long_section)
        guided_document["sections"] = tuple(current_sections)
        guided_document.update({
            "generation_mode": str(long_form.get("generation_mode") or "unknown"),
            "model_used": bool(long_form.get("model_used", False)),
            "review_verdict": "approved",
            "review_reason": str(long_form.get("review_reason") or ""),
            "target_characters": int(long_form.get("target_characters") or 0),
            "actual_characters": len(long_form_content),
            "source_passage_count": int(long_form.get("source_passage_count") or 0),
            "source_references": tuple(
                str(item)
                for item in long_form.get("source_references", ()) or ()
                if str(item)
            ),
            "retrieval_query_count": int(long_form.get("retrieval_query_count") or 0),
            "delivery_variant": str(long_form.get("delivery_variant") or "guided_long_read"),
            "collaboration_path": "Generation → Reviewer → Guidance",
        })

    diagnostic_first = bool(
        isinstance(teaching_decision, AdaptiveTeachingDecision)
        and teaching_decision.diagnostic_needed
    )
    research_first = bool(
        isinstance(teaching_decision, AdaptiveTeachingDecision)
        and (
            teaching_decision.content_depth in {"research", "advanced"}
            or "research_task" in resource_modes
        )
    )
    recommended_family = (
        ResourceFamily.ASSESSMENT
        if diagnostic_first
        else ResourceFamily.PRACTICAL
        if research_first
        else ResourceFamily.KNOWLEDGE
    )

    knowledge = LearningResource(
        resource_id=_resource_id(task_id, ResourceFamily.KNOWLEDGE.value, "guided_long_read"),
        resource_family=ResourceFamily.KNOWLEDGE,
        resource_form=("guided_long_read" if reviewed_answer else "prerequisite_or_gap_card"),
        title=(
            "基础学习讲义（待诊断）"
            if diagnostic_first
            else "个性化科研专题长文"
            if long_form_approved and research_first
            else "个性化概念讲义"
            if long_form_approved
            else "已审核学习讲义"
            if reviewed_answer
            else "发布前需要补足的概念与证据"
        ),
        target_concepts=target_concepts,
        prerequisite_concepts=prerequisite_concepts,
        learning_goal=goal,
        learner_fit_reason=(
            f"当前没有足够真实作答，先按 {depth} 深度呈现，并等待诊断结果。"
            if diagnostic_first
            else f"当前教学策略为 {strategy}，内容深度为 {depth}。"
        ),
        difficulty=depth,
        estimated_time_minutes=(
            max(12, min(30, round(len(long_form_content) / 240)))
            if long_form_approved
            else 12 if depth in {"foundation", "beginner"} else 18
        ),
        payload={
            "reviewed_summary": reviewed_answer,
            "concept_names": concept_names,
            "prerequisite_names": prerequisite_names,
            "knowledge_gap": knowledge_gaps,
            "guided_document": guided_document,
            "guided_questions": guided_questions,
            "guided_question_mode": (
                "llm_reviewed_socratic"
                if model_questions_approved
                else "relation_backed_fallback"
            ),
            "guided_question_review": (
                {
                    "model_used": True,
                    "reviewer_executed": True,
                    "review_verdict": "approved",
                    "review_reason": str(question_review.get("review_reason") or ""),
                    "source_passage_count": int(
                        question_review.get("source_passage_count") or 0
                    ),
                    "collaboration_path": str(
                        question_review.get("collaboration_path")
                        or "Generation → Reviewer → Guidance"
                    ),
                }
                if model_questions_approved
                else {
                    "model_used": False,
                    "reviewer_executed": False,
                    "review_verdict": "not_applicable",
                    "collaboration_path": "R06 Concept Relation → Guidance",
                }
            ),
            "recommended": recommended_family is ResourceFamily.KNOWLEDGE,
            "distribution_reason": (
                "当前需要先建立可审核的概念和机制理解。"
                if recommended_family is ResourceFamily.KNOWLEDGE
                else "当前优先项由 Diagnosis 的教学决策分发给其他资源形态。"
            ),
            "content_modes": tuple(dict.fromkeys((
                "guided_long_read",
                *(
                    teaching_decision.representation_modes
                    if isinstance(teaching_decision, AdaptiveTeachingDecision)
                    else ()
                ),
            ))),
        },
        source_type=(
            ResourceSourceType.GENERATED
            if long_form_approved and bool(long_form.get("model_used", False))
            else ResourceSourceType.DERIVED
        ),
        evidence_refs=evidence_refs,
        review_status=("approved" if reviewed_answer else "withheld"),
        provenance=(
            "learner_decision:diagnosis_interpreted",
            "knowledge:concept_context",
            "quality:release_gate",
            *(
                ("agent.knowledge.generation:long_form", "agent.quality.review:resource_review")
                if long_form_approved
                else ()
            ),
        ),
        interaction_actions=(
            ResourceInteractionAction.OPEN,
            ResourceInteractionAction.UNDERSTOOD,
            ResourceInteractionAction.STILL_CONFUSED,
            ResourceInteractionAction.CHANGE_EXPLANATION,
            ResourceInteractionAction.REQUEST_EXAMPLE,
            ResourceInteractionAction.DEEPEN,
            ResourceInteractionAction.ASK_FOLLOW_UP,
        ),
        completion_signal="explicit_feedback_only",
        next_action=("查看前置概念" if prerequisite_concepts else "继续理解当前概念"),
    )

    practical_steps: list[dict[str, Any]] = [{
        "name": "锁定本次任务",
        "operation": goal,
        "check": "后续记录必须回答这一任务，不扩展为无关主题。",
        "source": "task_goal",
    }]
    if concept_names:
        practical_steps.append({
            "name": "核对目标概念",
            "operation": "本任务已映射到：" + "、".join(concept_names),
            "check": (
                "待先补足：" + "、".join(prerequisite_names)
                if prerequisite_names
                else "当前没有发布额外先修缺口。"
            ),
            "source": "task_concept_mapping",
        })
    for claim in accepted_claims[:3]:
        refs = tuple(str(item) for item in claim.get("evidence_refs", ()) if str(item))
        practical_steps.append({
            "name": "核对已审核结论",
            "operation": str(claim.get("statement") or ""),
            "check": (
                "绑定证据：" + "、".join(refs)
                if refs
                else "本结论没有可公开的证据引用，不得外推。"
            ),
            "source": "reviewed_claim",
        })
    if knowledge_gaps:
        practical_steps.append({
            "name": "保留证据缺口",
            "operation": "当前未解决：" + "、".join(knowledge_gaps[:4]),
            "check": "缺失证据不能用模板参数或相邻材料结论补齐。",
            "source": "released_knowledge_gap",
        })
    practical_steps.append({
        "name": "输出证据分析记录",
        "operation": "按“结论—证据—条件—限制—下一步”整理本次学习或研究记录。",
        "check": "只能写入本任务已发布的事实；新推断必须单独标注。",
        "source": "quality_release",
    })
    practical_steps = [
        {"step": index, **step}
        for index, step in enumerate(practical_steps, start=1)
    ]
    practical = LearningResource(
        resource_id=_resource_id(task_id, ResourceFamily.PRACTICAL.value, "evidence_workbook"),
        resource_family=ResourceFamily.PRACTICAL,
        resource_form=(
            "research_task"
            if "research_task" in resource_modes
            else "evidence_analysis_workbook"
        ),
        title="当前任务证据分析工作单",
        target_concepts=target_concepts,
        prerequisite_concepts=prerequisite_concepts,
        learning_goal=goal,
        learner_fit_reason=f"当前策略需要把 {', '.join(concept_names) or '目标概念'} 转化为可验证任务。",
        difficulty=depth,
        estimated_time_minutes=max(12, min(30, len(practical_steps) * 4)),
        payload={
            "steps": tuple(practical_steps),
            "target_concept_names": concept_names,
            "prerequisite_names": prerequisite_names,
            "accepted_claim_count": len(accepted_claims),
            "knowledge_gap_count": len(knowledge_gaps),
            "task_binding": task_id,
            "parameter_status": "not_prescribed",
            "safety_boundary": "follow_local_lab_sop",
            "recommended": recommended_family is ResourceFamily.PRACTICAL,
            "distribution_reason": (
                "当前为科研深入或实践导向，优先分发可验证任务。"
                if recommended_family is ResourceFamily.PRACTICAL
                else "当前不是首要形态，但保留为可选科研活动。"
            ),
        },
        source_type=ResourceSourceType.DERIVED,
        evidence_refs=evidence_refs,
        review_status="derived_from_released_task",
        provenance=(
            "task:current",
            "knowledge:concept_context",
            "quality:release_gate",
            "collaboration:accepted_claims",
        ),
        interaction_actions=(
            ResourceInteractionAction.OPEN,
            ResourceInteractionAction.START_PRACTICE,
            ResourceInteractionAction.NEXT_CONCEPT,
        ),
        completion_signal="user_action_or_assessment_required",
        next_action="使用本任务的真实结论与证据完成分析记录",
    )

    target_kps: tuple[str, ...] = ()
    if isinstance(knowledge_context, KnowledgeLearningContext):
        target_kps = tuple(dict.fromkeys(
            kp
            for concept_id in (*prerequisite_concepts, *target_concepts)
            for kp in knowledge_context.concept_to_kps.get(concept_id, ())
        ))
    assessment = LearningResource(
        resource_id=_resource_id(task_id, ResourceFamily.ASSESSMENT.value, "practice_bank_launch"),
        resource_family=ResourceFamily.ASSESSMENT,
        resource_form="practice_bank_launch",
        title="基于现有题库的分阶练习",
        target_concepts=target_concepts,
        prerequisite_concepts=prerequisite_concepts,
        learning_goal=goal,
        learner_fit_reason=(
            "真实作答会形成 AnswerRecord 并进入现有 BKT/IRT；仅阅读不会改变掌握度。"
        ),
        difficulty=(
            teaching_decision.difficulty_strategy
            if isinstance(teaching_decision, AdaptiveTeachingDecision)
            else "diagnose_then_maintain"
        ),
        estimated_time_minutes=10,
        payload={
            "endpoint": "/l2/practice/questions",
            "target_kps": target_kps,
            "question_source": "local_practice_bank",
            "availability_policy": "empty_when_no_authored_question",
            "stages": (
                {
                    "stage": "diagnostic",
                    "label": "诊断",
                    "attempt_purpose": "DIAGNOSTIC",
                    "use": "先获得真实作答证据，不预设学习者强弱。",
                },
                {
                    "stage": "consolidation",
                    "label": "巩固",
                    "attempt_purpose": "REQUIRED_PRACTICE",
                    "use": "围绕当前 Concept 的已编写题目进行练习。",
                },
                {
                    "stage": "challenge",
                    "label": "进阶",
                    "attempt_purpose": "STAGED_ASSESSMENT",
                    "use": "仅由当前难度决策和题库实际覆盖决定是否采用。",
                },
            ),
            "stage_selection": (
                "challenge"
                if isinstance(teaching_decision, AdaptiveTeachingDecision)
                and teaching_decision.difficulty_strategy in {"increase", "raise"}
                else "diagnostic"
                if isinstance(teaching_decision, AdaptiveTeachingDecision)
                and teaching_decision.diagnostic_needed
                else "consolidation"
            ),
            "recommended": recommended_family is ResourceFamily.ASSESSMENT,
            "distribution_reason": (
                "当前学习者证据不足，先用真实作答建立诊断事实。"
                if recommended_family is ResourceFamily.ASSESSMENT
                else "练习仅在提交真实作答后更新学习模型。"
            ),
        },
        source_type=ResourceSourceType.RETRIEVED,
        evidence_refs=(),
        review_status="authored_local_questions",
        provenance=("L2:PracticeBank", "L2:AnswerRecord", "L2:BKT/IRT"),
        interaction_actions=(ResourceInteractionAction.START_PRACTICE,),
        completion_signal="submitted_answer_record",
        next_action="完成真实作答后更新学习模型",
    )
    return LearningResourcePlan(
        task_id=task_id,
        learner_id=learner_id,
        teaching_depth=depth,
        strategy=strategy,
        resources=(knowledge, practical, assessment),
    )


def public_resource_projection(plan: LearningResourcePlan) -> list[dict[str, Any]]:
    return [
        {
            "resource_id": item.resource_id,
            "resource_family": item.resource_family.value,
            "resource_form": item.resource_form,
            "title": item.title,
            "target_concepts": list(item.target_concepts),
            "prerequisite_concepts": list(item.prerequisite_concepts),
            "learning_goal": item.learning_goal,
            "learner_fit_reason": item.learner_fit_reason,
            "difficulty": item.difficulty,
            "estimated_time_minutes": item.estimated_time_minutes,
            "payload": dict(item.payload),
            "source_type": item.source_type.value,
            "evidence_refs": list(item.evidence_refs),
            "review_status": item.review_status,
            "provenance": list(item.provenance),
            "interaction_actions": [action.value for action in item.interaction_actions],
            "completion_signal": item.completion_signal,
            "next_action": item.next_action,
        }
        for item in plan.resources
    ]


def render_learning_resource_markdown(resource: Mapping[str, Any]) -> str:
    """Render one learning resource into a downloadable Markdown document.

    Assembles only already-reviewed content (title/goal/reviewed answer/
    guided_document sections/concepts/gaps/evidence).  It creates no new
    scientific claim; it only serializes the released resource into long-form
    Markdown so the front end can present and download it as a document.
    """
    title = str(resource.get("title") or "学习讲义").strip()
    goal = str(resource.get("learning_goal") or "").strip()
    fit = str(resource.get("learner_fit_reason") or "").strip()
    depth = str(resource.get("difficulty") or "").strip()
    review_status = str(resource.get("review_status") or "").strip()
    concepts = [str(c).strip() for c in (resource.get("target_concepts") or ()) if str(c).strip()]
    prereqs = [str(c).strip() for c in (resource.get("prerequisite_concepts") or ()) if str(c).strip()]
    evidence = [str(e).strip() for e in (resource.get("evidence_refs") or ()) if str(e).strip()]
    payload = resource.get("payload") if isinstance(resource.get("payload"), Mapping) else {}
    reviewed_summary = str(payload.get("reviewed_summary") or "").strip()
    gaps = [str(g).strip() for g in (payload.get("knowledge_gap") or ()) if str(g).strip()]
    guided = payload.get("guided_document") if isinstance(payload.get("guided_document"), Mapping) else {}
    sections = guided.get("sections") or ()

    lines: list[str] = [f"# {title}", ""]
    if goal:
        lines += [f"> **学习目标**：{goal}", ""]
    if fit:
        lines += [f"> **适配说明**：{fit}", ""]
    lines += [f"- 难度：{depth}", f"- 审核状态：{review_status}", ""]
    if concepts:
        lines += ["## 目标概念", "", "、".join(concepts), ""]
    if prereqs:
        lines += ["## 前置概念", "", "、".join(prereqs), ""]
    if reviewed_summary:
        lines += ["## 已审核答案", "", reviewed_summary, ""]
    if sections:
        lines += ["## 讲义正文", ""]
        for section in sections:
            if not isinstance(section, Mapping):
                continue
            section_title = str(section.get("title") or "小节").strip()
            content = str(section.get("content") or "").strip()
            if not content:
                continue
            lines += [f"### {section_title}", "", content, ""]
    if gaps:
        lines += ["## 知识缺口", "", *[f"- {g}" for g in gaps], ""]
    if evidence:
        lines += ["## 证据来源", "", *[f"- {e}" for e in evidence], ""]
    lines += ["---", "", "*本讲义由系统依据已审核证据自动生成，可溯源、可下载。*"]
    return "\n".join(lines)


def render_learning_resource_docx(resource: Mapping[str, Any]) -> bytes:
    """Render a learning resource into a downloadable Word (.docx) document."""
    import io

    from docx import Document

    md = render_learning_resource_markdown(resource)
    doc = Document()
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("> "):
            doc.add_paragraph(line[2:], style="Intense Quote")
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line == "---":
            doc.add_paragraph("_" * 30)
        elif line.startswith("*") and line.endswith("*"):
            doc.add_paragraph(line.strip("*"))
        else:
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_learning_resource_pdf(resource: Mapping[str, Any]) -> bytes:
    """Render a learning resource into a downloadable PDF document."""
    import io
    import os

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    # 注册中文字体 (黑体/宋体/微软雅黑 多级回退)
    _font_candidates = (
        (r"C:\Windows\Fonts\simhei.ttf", 0),
        (r"C:\Windows\Fonts\msyh.ttc", 0),
        (r"C:\Windows\Fonts\simsun.ttc", 0),
    )
    _font_name = None
    for _path, _sub in _font_candidates:
        if os.path.exists(_path):
            try:
                pdfmetrics.registerFont(TTFont("CN", _path, subfontIndex=_sub))
                _font_name = "CN"
                break
            except Exception:  # noqa: BLE001
                continue
    if _font_name is None:
        # 无中文字体时退回内置 Helvetica (中文将显示为方块, 但流程不中断)
        _font_name = "Helvetica"

    md = render_learning_resource_markdown(resource)
    body = ParagraphStyle("CN", fontName=_font_name, fontSize=11, leading=18)
    h1 = ParagraphStyle("H1", parent=body, fontSize=18, leading=24, spaceAfter=8)
    h2 = ParagraphStyle("H2", parent=body, fontSize=14, leading=20, spaceBefore=8, spaceAfter=4)
    h3 = ParagraphStyle("H3", parent=body, fontSize=12, leading=18, spaceBefore=6, spaceAfter=3)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    story = []
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            story.append(Paragraph(line[4:], h3))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], h2))
        elif line.startswith("# "):
            story.append(Paragraph(line[2:], h1))
        elif line.startswith("> "):
            story.append(Paragraph(line[2:], body))
        elif line.startswith("- "):
            story.append(Paragraph("• " + line[2:], body))
        elif line == "---":
            story.append(Spacer(1, 10))
        elif line.startswith("*") and line.endswith("*"):
            story.append(Paragraph(line.strip("*"), body))
        else:
            story.append(Paragraph(line, body))
    doc.build(story)
    return buf.getvalue()




def build_resource_interaction_event(
    *,
    learner_id: str,
    task_id: str,
    resource: Mapping[str, Any],
    action: str,
) -> ResourceInteractionEvent:
    parsed_action = ResourceInteractionAction(str(action))
    allowed = set(str(item) for item in resource.get("interaction_actions") or ())
    if parsed_action.value not in allowed:
        raise ValueError("interaction action is not allowed for this resource")
    resource_id = str(resource.get("resource_id") or "")
    if not learner_id or not task_id or not resource_id:
        raise ValueError("learner_id, task_id and resource_id are required")
    timestamp = time.time()
    seed = f"{learner_id}|{task_id}|{resource_id}|{parsed_action.value}|{timestamp}"
    return ResourceInteractionEvent(
        event_id=f"resource-event-{sha256(seed.encode('utf-8')).hexdigest()[:20]}",
        learner_id=learner_id,
        task_id=task_id,
        resource_id=resource_id,
        resource_family=str(resource.get("resource_family") or ""),
        resource_form=str(resource.get("resource_form") or ""),
        action=parsed_action,
        concept_ids=tuple(str(item) for item in resource.get("target_concepts") or () if str(item)),
        source_type=str(resource.get("source_type") or "unknown"),
        timestamp=timestamp,
    )


__all__ = [
    "LearningResource",
    "LearningResourcePlan",
    "ResourceFamily",
    "ResourceInteractionAction",
    "ResourceInteractionEvent",
    "ResourceSourceType",
    "build_learning_resource_plan",
    "build_resource_interaction_event",
    "public_resource_projection",
]
